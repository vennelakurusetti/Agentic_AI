"""
create_doc.py - Scrape, clean, and create a structured Word document knowledge base
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Import from sibling modules
from scraper import scrape_all_pages
from clean_data import clean_text


# Structured headings for the knowledge base document
STRUCTURED_SECTIONS = [
    ("1. About BVRIT Hyderabad", [
        "About BVRIT Hyderabad",
        "About Society – SVES",
    ]),
    ("2. Vision & Mission", [
        "Vision & Mission",
    ]),
    ("3. Management & Leadership", [
        "Management",
        "Principal",
        "Organogram",
    ]),
    ("4. Accreditations & Rankings", [
        "NIRF",
    ]),
    ("5. Departments – UG Programs", [
        "CSE Department",
        "CSE (AI&ML) Department",
        "ECE Department",
        "EEE Department",
        "IT Department",
        "BS&H Department",
        "Honor Degree Program",
    ]),
    ("6. Admissions", [
        "Admission Process",
        "EAMCET Ranks",
        "B-Category Admission",
        "Fee Details",
        "Intake of Courses",
        "Documents to Submit",
        "Hostel Admission",
        "Transportation",
        "PM Vidyalaxmi Scheme",
    ]),
    ("7. Placements", [
        "Placements – Training & Placement Process",
        "Placements – Training & Placement Cell",
        "Placements – Team",
        "Placements – Employability Skills",
        "Placements – Internships",
        "Placements – Placement Details",
        "Placements – Testimonials",
    ]),
    ("8. Campus Facilities", [
        "Library",
        "Food and Cafetaria",
        "Gym",
        "Temple",
        "Security",
        "PCS Facilities",
        "Entry – Exit System",
    ]),
    ("9. Research & Development", [
        "Research & Development",
        "Faculty Thrust Areas",
        "Ph.D Awarded",
        "Publications",
        "Research Projects",
        "Consultancy Projects",
        "Patents Published",
        "Patents Granted",
        "Center of Excellence (CoEs)",
    ]),
    ("10. Differentiators & Special Centers", [
        "VSSC (Vishnu Student Success Center)",
        "GSAC (Graduate Study Abroad Center)",
        "IIC BVRITH",
        "EDC (Entrepreneurship Development Cell)",
        "Student Affairs Council (SAC)",
        "IKS and NEP",
    ]),
    ("11. Student Activities & Clubs", [
        "NSS Club",
        "Rotaract Club",
        "Sports Club",
        "CSI",
        "IEEE",
        "ACM",
    ]),
    ("12. Alumni", [
        "Alumni Speak",
    ]),
    ("13. Contact Details", [
        "Contact Us",
    ]),
]


def add_heading_styled(doc, text, level=1):
    """Add a styled heading to the document."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)
    return heading


def add_body_text(doc, text):
    """Add cleaned body text to the document."""
    paragraphs = text.split("\n")
    for para_text in paragraphs:
        para_text = para_text.strip()
        if para_text:
            p = doc.add_paragraph(para_text)
            p.style.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15


def create_knowledge_base(data):
    """Create a structured Word document knowledge base."""
    doc = Document()
    
    # ---- Title Page ----
    title = doc.add_heading("BVRIT Hyderabad College of Engineering for Women", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)
    
    subtitle = doc.add_paragraph("College FAQ Knowledge Base")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(102, 102, 102)
    
    doc.add_paragraph("")  # spacing
    info = doc.add_paragraph("Generated from the official website: https://bvrithyderabad.edu.in")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.runs[0].font.size = Pt(10)
    info.runs[0].font.italic = True
    
    doc.add_page_break()
    
    # ---- Table of Contents ----
    toc_heading = doc.add_heading("Table of Contents", level=1)
    for run in toc_heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)
    
    for section_title, _ in STRUCTURED_SECTIONS:
        toc_p = doc.add_paragraph(section_title, style="List Number")
        toc_p.paragraph_format.space_after = Pt(2)
    
    doc.add_page_break()
    
    # ---- Content Sections ----
    for section_title, page_keys in STRUCTURED_SECTIONS:
        add_heading_styled(doc, section_title, level=1)
        
        for key in page_keys:
            if key in data and data[key].strip():
                # Only add content if it's not an error message
                if not data[key].startswith("Error scraping"):
                    add_heading_styled(doc, key, level=2)
                    add_body_text(doc, data[key])
                    doc.add_paragraph("")  # spacing
        
        doc.add_page_break()
    
    # ---- Save ----
    output_path = os.path.join("output", "BVRIT_Hyderabad_Knowledge_Base.docx")
    doc.save(output_path)
    print(f"\nDocument saved to: {output_path}")
    return output_path


def main():
    """Main function: scrape -> clean -> create doc."""
    print("=" * 60)
    print("BVRIT Hyderabad Knowledge Base Generator")
    print("=" * 60)
    
    # Step 1: Scrape
    print("\n[Step 1/3] Scraping all pages...")
    raw_data = scrape_all_pages()
    
    # Step 2: Clean
    print("\n[Step 2/3] Cleaning extracted text...")
    cleaned_data = {}
    for title, text in raw_data.items():
        print(f"  Cleaning: {title} ... ", end="", flush=True)
        cleaned_data[title] = clean_text(text)
        print(f"done ({len(cleaned_data[title])} chars)")
    
    # Step 3: Create Document
    print("\n[Step 3/3] Creating structured Word document...")
    output_path = create_knowledge_base(cleaned_data)
    
    print(f"\n✅ Knowledge base created successfully!")
    print(f"   File: {output_path}")
    print(f"   Total sections: {len(STRUCTURED_SECTIONS)}")
    print(f"   Total pages: {len(cleaned_data)}")


if __name__ == "__main__":
    main()
