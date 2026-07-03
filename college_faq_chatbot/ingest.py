"""
ingest.py - Document ingestion pipeline for the College FAQ Chatbot.
Uses python-docx to parse DOCX preserving heading styles as section metadata.
"""

import os
import re
import shutil
from pathlib import Path
from typing import List, Optional, Tuple

from dotenv import load_dotenv
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document

import config
from utils import logger, Timer

load_dotenv()


# Track whether we've auto-rebuilt (persist check via flag file)
_AUTO_REBUILT_FLAG = config.CHROMA_DIR / ".auto_rebuilt"


def parse_docx_paragraphs(docx_path: Path) -> List[Document]:
    """
    Read a DOCX file paragraph by paragraph using python-docx.
    Detects Heading 1 and Heading 2 styles and attaches section metadata.
    """
    if not docx_path.exists():
        raise FileNotFoundError(f"Document not found: {docx_path}")

    logger.info(f"Loading document with python-docx: {docx_path}")
    doc = DocxDocument(str(docx_path))

    paragraphs_data: List[Tuple[str, str]] = []  # (section, text)
    current_section = "General"
    heading_stack = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ""

        # Detect heading styles
        is_heading = False
        if "Heading 1" in style_name or style_name == "Heading 1":
            current_section = text
            heading_stack = [text]
            is_heading = True
        elif "Heading 2" in style_name or style_name == "Heading 2":
            if heading_stack:
                current_section = f"{heading_stack[0]} - {text}"
            else:
                current_section = text
            heading_stack = heading_stack[:1] + [text]
            is_heading = True
        elif "heading" in style_name.lower() or style_name.startswith("Heading"):
            if heading_stack:
                current_section = f"{heading_stack[0]} - {text}"
            else:
                current_section = text
            heading_stack.append(text)
            is_heading = True

        # Also detect numbered headings like "1. About BVRIT Hyderabad"
        if not is_heading:
            match = re.match(r'^(\d+)\.\s+(.+)', text)
            if match:
                current_section = text
                heading_stack = [text]
                is_heading = True

        # Detect ALL CAPS short headings
        if not is_heading and len(text) > 3 and len(text) < 60 and text.isupper():
            current_section = text
            heading_stack = [text]
            is_heading = True

        if not is_heading:
            paragraphs_data.append((current_section, text))

    logger.info(f"Parsed {len(paragraphs_data)} paragraphs from {len(doc.paragraphs)} total")
    sections_found = set(s for s, _ in paragraphs_data)
    logger.info(f"Sections detected: {sorted(sections_found)}")

    # Group paragraphs by section for better chunking
    section_groups: dict = {}
    for section, text in paragraphs_data:
        if section not in section_groups:
            section_groups[section] = []
        section_groups[section].append(text)

    documents = []
    for section, texts in section_groups.items():
        combined = "\n".join(texts)
        doc = Document(
            page_content=combined,
            metadata={
                "section": section,
                "source": config.DOCX_PATH.name,
            }
        )
        documents.append(doc)

    logger.info(f"Created {len(documents)} section-grouped documents")
    return documents


def split_documents(documents: List[Document]) -> List[Document]:
    """Split documents into chunks, preserving section metadata."""
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
        length_function=len,
    )

    all_chunks = []
    for doc in documents:
        section = doc.metadata.get("section", "General")
        chunks = text_splitter.split_text(doc.page_content)
        for i, chunk_text in enumerate(chunks):
            chunk_doc = Document(
                page_content=chunk_text,
                metadata={
                    "section": section,
                    "source": config.DOCX_PATH.name,
                    "chunk_index": len(all_chunks) + i,
                }
            )
            all_chunks.append(chunk_doc)

    logger.info(
        f"Split into {len(all_chunks)} chunks "
        f"(size={config.CHUNK_SIZE}, overlap={config.CHUNK_OVERLAP})"
    )

    # Log sections present
    sections_in_chunks = set(c.metadata["section"] for c in all_chunks)
    logger.info(f"Sections in chunks: {sorted(sections_in_chunks)}")

    return all_chunks


def get_embeddings() -> OpenAIEmbeddings:
    """Get the OpenAI embeddings model configured for OpenRouter."""
    api_key = os.getenv("OPENROUTER_API_KEY")
    if not api_key:
        raise ValueError("OPENROUTER_API_KEY not found in environment variables")

    return OpenAIEmbeddings(
        model=config.EMBEDDING_MODEL,
        openai_api_key=api_key,
        openai_api_base=config.OPENROUTER_BASE_URL,
    )


def needs_rebuild() -> bool:
    """
    Check if the vector store needs to be rebuilt.
    Returns True if no .auto_rebuilt flag exists (first run after this update).
    """
    if not config.CHROMA_DIR.exists() or not any(config.CHROMA_DIR.iterdir()):
        return True
    if not _AUTO_REBUILT_FLAG.exists():
        return True
    return False


def mark_rebuilt() -> None:
    """Create the auto-rebuilt flag file."""
    _AUTO_REBUILT_FLAG.touch()
    logger.info("Marked vector store as rebuilt with new metadata")


def create_vector_store(
    chunks: List[Document],
    force_recreate: bool = False,
) -> Chroma:
    """Create or load a persistent ChromaDB vector store."""
    chroma_path = str(config.CHROMA_DIR)

    if force_recreate and config.CHROMA_DIR.exists():
        shutil.rmtree(chroma_path)
        logger.info("Removed existing ChromaDB due to force_recreate=True")

    # Auto-rebuild if needed (first run after metadata upgrade)
    if needs_rebuild() and not force_recreate:
        logger.info("Auto-rebuild needed: old metadata detected. Rebuilding vector store...")
        if config.CHROMA_DIR.exists():
            shutil.rmtree(chroma_path)
        force_recreate = True

    embeddings = get_embeddings()

    if config.CHROMA_DIR.exists() and any(config.CHROMA_DIR.iterdir()) and not force_recreate:
        logger.info(f"Loading existing vector store from {chroma_path}")
        vector_store = Chroma(
            collection_name=config.CHROMA_COLLECTION_NAME,
            embedding_function=embeddings,
            persist_directory=chroma_path,
        )
        logger.info("Loaded vector store with existing documents")
    else:
        logger.info(f"Creating new vector store at {chroma_path}")
        vector_store = Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            collection_name=config.CHROMA_COLLECTION_NAME,
            persist_directory=chroma_path,
        )
        logger.info(f"Created vector store with {len(chunks)} chunks")
        mark_rebuilt()

    return vector_store


def ingest_document(force_recreate: bool = False) -> Chroma:
    """Run the full ingestion pipeline: parse, split, embed, store."""
    with Timer("Document ingestion"):
        documents = parse_docx_paragraphs(config.DOCX_PATH)
        chunks = split_documents(documents)
        vector_store = create_vector_store(chunks, force_recreate=force_recreate)

    logger.info("Ingestion complete!")
    return vector_store


def get_chunk_count() -> Optional[int]:
    """Get the number of chunks in the vector store (without recreating)."""
    if config.CHROMA_DIR.exists() and any(config.CHROMA_DIR.iterdir()):
        try:
            embeddings = get_embeddings()
            vector_store = Chroma(
                collection_name=config.CHROMA_COLLECTION_NAME,
                embedding_function=embeddings,
                persist_directory=str(config.CHROMA_DIR),
            )
            return vector_store._collection.count()
        except Exception as e:
            logger.warning(f"Could not get chunk count: {e}")
    return None


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Ingest college knowledge base document")
    parser.add_argument("--force", action="store_true", help="Force recreate vector store")
    args = parser.parse_args()

    vector_store = ingest_document(force_recreate=args.force)
    count = vector_store._collection.count()
    print(f"\n✅ Ingestion complete! {count} chunks stored in ChromaDB.")